import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

DOCX_DIR = "/Users/rajchoudhary/YuvaIntern-Content-Writer-Internship/reports_docx"
MD_DIR = "/Users/rajchoudhary/YuvaIntern-Content-Writer-Internship/reports_markdown"
os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(MD_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document_base(title_text, subtitle_text, week_str, word_count_str):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Table (Header Card)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "0F1E36")
    set_cell_margins(cell, top=200, bottom=200, left=250, right=250)
    
    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run("YUVAINTERN VIRTUAL INTERNSHIP PROGRAM  •  CONTENT WRITING")
    r0.font.name = "Calibri"
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(56, 189, 248) # Cyan
    
    p1 = cell.add_paragraph()
    r1 = p1.add_run(title_text)
    r1.font.name = "Georgia"
    r1.font.size = Pt(17)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = cell.add_paragraph()
    r2 = p2.add_run(f"Candidate: Raj Choudhary (craj40215@gmail.com)  |  {week_str}  |  Estimated Length: {word_count_str}")
    r2.font.name = "Calibri"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(203, 213, 225)
    
    doc.add_paragraph() # Spacer
    return doc

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 30, 54) # Deep Navy
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(2, 132, 199) # Brand Blue
    return h

def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    return p

def add_bullet_point(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_bold = p.add_run(bold_prefix + " ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(11)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(15, 23, 42)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout_box(doc, title, text, bg_hex="F0F9FF", border_color="0284C7"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    r_title = p.add_run(f"★  {title.upper()}\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(2, 132, 199)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph() # Spacer

print("Helper base initialized.")
